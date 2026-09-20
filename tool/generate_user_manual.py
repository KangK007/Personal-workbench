from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "manual" / "USER_MANUAL.md"
OUTPUT = ROOT / "docs" / "manual" / "个人工作台_用户使用手册.docx"
APP_VERSION = "0.2.1+6"

GREEN = "2F7D57"
DARK_GREEN = "1F4D3A"
INK = "1B2B20"
MUTED = "617268"
TABLE_FILL = "E8F2EC"
BORDER = "B9CEC0"


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_fixed_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def next_numbering_id(numbering) -> tuple[int, int]:
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    return (max(abstract_ids, default=-1) + 1, max(num_ids, default=0) + 1)


def create_numbering(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id, num_id = next_numbering_id(numbering)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if kind == "number" else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if kind == "number" else "•")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_numbered_paragraph(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    add_inline(paragraph, text)


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, color=INK, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", east_asia="Microsoft YaHei", size=9.5, color=DARK_GREEN)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, color=INK)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    column_count = len(rows[0])
    if column_count == 2:
        widths = [3000, 6360]
    elif column_count == 3:
        widths = [2200, 3580, 3580]
    elif column_count == 4:
        widths = [1600, 2600, 2580, 2580]
    else:
        base = 9360 // column_count
        widths = [base] * column_count
        widths[-1] += 9360 - sum(widths)

    table = doc.add_table(rows=len(rows), cols=column_count)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, color=INK, bold=row_index == 0)
            if row_index == 0:
                set_cell_fill(cell, TABLE_FILL)
    set_fixed_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_picture(doc: Document, alt: str, image_path: Path) -> None:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    max_width = 6.25
    max_height = 5.9
    width = min(max_width, max_height * ratio)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", alt)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, GREEN, 18, 10),
        ("Heading 2", 13, GREEN, 14, 7),
        ("Heading 3", 12, DARK_GREEN, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_together = True


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = paragraph.add_run("个人工作台用户手册")
    set_run_font(left, size=9, color=MUTED, bold=True)
    right = paragraph.add_run(f"\t{APP_VERSION}")
    set_run_font(right, size=9, color=MUTED)
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("第 ")
    set_run_font(label, size=9, color=MUTED)
    add_field(paragraph, "PAGE")
    label = paragraph.add_run(" 页")
    set_run_font(label, size=9, color=MUTED)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(92)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("个人工作台")
    set_run_font(run, size=30, color=DARK_GREEN, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("用户使用手册")
    set_run_font(run, size=16, color=GREEN)
    for text, size, bold in (
        ("Windows 与 Android", 11, True),
        (f"版本 {APP_VERSION}", 10, False),
        ("依据实际代码、完整测试矩阵和真实运行界面编制", 10, False),
        ("验证日期：2026-09-19", 10, False),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(text)
        set_run_font(run, size=size, color=MUTED, bold=bold)
    doc.add_page_break()


def add_contents(doc: Document, headings: list[str]) -> None:
    title = doc.add_paragraph("目录", style="Heading 1")
    title.paragraph_format.space_before = Pt(0)
    for heading in headings:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(heading)
        set_run_font(run, size=10.5, color=INK)
    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def build_document() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "个人工作台用户使用手册"
    doc.core_properties.subject = f"个人工作台 {APP_VERSION} Windows 与 Android 使用说明"
    doc.core_properties.author = "个人工作台项目"
    doc.core_properties.keywords = "个人工作台, Windows, Android, 任务, 项目, 专注, 回顾"
    add_cover(doc)
    add_contents(doc, headings)

    index = 0
    active_list: str | None = None
    active_num_id: int | None = None
    image_count = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped or stripped.startswith("# 个人工作台用户使用手册"):
            active_list = None
            index += 1
            continue
        if re.match(r"^(版本|适用平台|手册验证日期)：", stripped):
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_paragraph(stripped[3:].strip(), style="Heading 1")
            active_list = None
        elif stripped.startswith("### "):
            doc.add_paragraph(stripped[4:].strip(), style="Heading 2")
            active_list = None
        elif re.match(r"^\d+\.\s+", stripped):
            if active_list != "number":
                active_list = "number"
                active_num_id = create_numbering(doc, "number")
            add_numbered_paragraph(doc, re.sub(r"^\d+\.\s+", "", stripped), active_num_id)
        elif stripped.startswith("- "):
            if active_list != "bullet":
                active_list = "bullet"
                active_num_id = create_numbering(doc, "bullet")
            add_numbered_paragraph(doc, stripped[2:].strip(), active_num_id)
        elif stripped.startswith("|"):
            rows, next_index = parse_table(lines, index)
            add_table(doc, rows)
            index = next_index - 1
            active_list = None
        else:
            image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
            if image_match:
                alt, relative = image_match.groups()
                image_path = (SOURCE.parent / relative).resolve()
                if not image_path.is_file():
                    raise FileNotFoundError(image_path)
                add_picture(doc, alt, image_path)
                image_count += 1
            else:
                style = "Caption" if re.match(r"^图\s*\d+", stripped) else None
                paragraph = doc.add_paragraph(style=style)
                if style == "Caption":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, stripped)
            active_list = None
        index += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    if image_count != 21:
        raise RuntimeError(f"Expected 21 screenshots, found {image_count}")

    with ZipFile(OUTPUT) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        numbering_xml = archive.read("word/numbering.xml").decode("utf-8")
        if document_xml.count("<w:tbl>") == 0:
            raise RuntimeError("Expected fixed-width tables")
        if document_xml.count("<w:tblHeader") != len(doc.tables):
            raise RuntimeError("Table header accessibility audit failed")
        if 'w:w="9360"' not in document_xml or 'w:type="fixed"' not in document_xml:
            raise RuntimeError("Table geometry audit failed")
        if 'w:line="300"' not in styles_xml:
            raise RuntimeError("Body 1.25 line spacing is missing")
        if 'w:left="540"' not in numbering_xml or 'w:hanging="270"' not in numbering_xml:
            raise RuntimeError("Numbering geometry audit failed")
        if "PLACEHOLDER" in document_xml:
            raise RuntimeError("Placeholder text remains")

    print(f"Created {OUTPUT}")
    print(f"Screenshots embedded: {image_count}")
    print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}")


if __name__ == "__main__":
    build_document()
