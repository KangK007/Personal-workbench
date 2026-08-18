from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "USER_GUIDE.md"
OUTPUT = ROOT / "docs" / "个人工作台Windows版新手使用指导.docx"
IMAGE_ROOT = ROOT / "test" / "goldens"
DOC_IMAGE_ROOT = ROOT / "docs" / "images"

FONT = "Microsoft YaHei"
BODY_COLOR = "252525"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "6B7280"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "4E8A59"
ORANGE = "C86848"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths_dxa):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_border(paragraph, color: str = "D7DBE2", bottom: bool = False, left: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    side = "bottom" if bottom else "left"
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "8" if bottom else "16")
    border.set(qn("w:space"), "6")
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_field(paragraph, field: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    begin_run = paragraph.add_run()._r
    begin_run.append(begin)
    instr_run = paragraph.add_run()._r
    instr_run.append(instr)
    separate_run = paragraph.add_run()._r
    separate_run.append(separate)
    value_run = paragraph.add_run()._r
    value_run.append(text)
    end_run = paragraph.add_run()._r
    end_run.append(end)


def add_image_alt(inline_shape, text: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", text)
    doc_pr.set("title", text)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY_COLOR)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    styles = {
        "Title": (28, DARK_BLUE, 0, 10),
        "Subtitle": (13, MUTED, 0, 14),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (11.5, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in styles.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.12
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = doc.styles.add_style("Guide Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(9.2)
    code.font.color.rgb = RGBColor.from_string("334155")
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.left_indent = Inches(0.22)
    code.paragraph_format.right_indent = Inches(0.22)
    code.paragraph_format.line_spacing = 1.12

    caption = doc.styles.add_style("Guide Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.12


def setup_page_furniture(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("个人工作台  |  Windows 版新手使用指导")
    set_run_font(run, size=8.5, color=MUTED)
    set_paragraph_border(p, color="D7DBE2", bottom=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(p, "PAGE")
    run = p.add_run(" 页")
    set_run_font(run, size=8.5, color=MUTED)


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(76)
    icon = ROOT / "assets" / "branding" / "app_icon.png"
    if icon.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(str(icon), width=Inches(0.92))
        add_image_alt(shape, "个人工作台应用图标")
        p.paragraph_format.space_after = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("个人工作台\nWindows 版新手使用指导")
    set_run_font(run, size=27, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("八栏工作台逐项说明 · 从第一次启动到每日收尾")
    set_run_font(run, size=13.5, color=MUTED)

    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [2340, 2340, 2340, 2340])
    values = [
        ("01", "今日", "执行与每日收尾"),
        ("02", "任务", "周期、周视图与任务群"),
        ("03", "项目", "任务、里程碑与资料"),
        ("04", "专注", "预设、计时与应用检测"),
        ("05", "笔记", "Markdown 与图片附件"),
        ("06", "回顾", "日周月事实快照"),
        ("07", "国策", "RSIP 树与轮次历史"),
        ("08", "设置", "提醒、托盘与数据安全"),
    ]
    for cell, (number, title_text, detail) in zip(
        [cell for row in table.rows for cell in row.cells], values
    ):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(number)
        set_run_font(run, size=15, color=ORANGE, bold=True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title_text)
        set_run_font(run, size=11, color=DARK_BLUE, bold=True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(detail)
        set_run_font(run, size=8.5, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("适用版本：当前 Flutter Windows / Android 本地优先版本")
    set_run_font(run, size=9.5, color=MUTED)
    doc.add_page_break()


def add_reading_map(doc: Document) -> None:
    heading = doc.add_paragraph("先看这里：第一次使用路线", style="Heading 1")
    set_keep_with_next(heading)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("第一次使用无需一次学完所有功能。先按下面顺序完成一次闭环，再按需要阅读周期任务、专注检测、国策和数据管理章节。")
    set_run_font(run)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [1300, 3000, 5060])
    headers = ["顺序", "先做什么", "对应章节"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("1", "认识八栏和全局入口", "第 1 章"),
        ("2", "创建今天的第一项任务", "第 2、3 章"),
        ("3", "开始专注并结算结果", "第 3、6 章"),
        ("4", "完成每日收尾并保存日回顾", "第 3、8 章"),
        ("5", "需要时再学习周期、任务链和国策", "第 4、9、10 章"),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row_values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, size=9.5, color=BODY_COLOR)

    add_note_box(doc, "阅读提示", "本文中的界面图来自当前项目的真实 Windows 测试界面。图中的示例任务和时间只用于说明操作，实际数据以你的本地记录为准。")


def add_note_box(doc: Document, label: str, text: str, color: str = CALLOUT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_repeat_table_header(table.rows[0])
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label)
    set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=9.5, color=BODY_COLOR)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, image_path: Path, caption_text: str, width_in: float) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f"Missing guide image: {image_path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    shape = p.add_run().add_picture(str(image_path), width=Inches(width_in))
    add_image_alt(shape, caption_text)
    caption = doc.add_paragraph(style="Guide Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(caption_text)
    set_run_font(run, size=9, color=MUTED)


def add_inline_text(paragraph, text: str, size: float | None = None) -> None:
    position = 0
    for match in re.finditer(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=DARK_BLUE)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size)


def split_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def parse_markdown(doc: Document, lines: list[str]) -> None:
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index].rstrip()

        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph(style="Guide Code")
                for code_line_index, code_line in enumerate(code_lines):
                    run = p.add_run(code_line)
                    set_run_font(run, size=9.2, color="334155")
                    run.font.name = "Consolas"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
                    if code_line_index < len(code_lines) - 1:
                        run.add_break()
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        image_match = re.match(r"^!\[(.+)]\((.+)\)$", line)
        if image_match:
            image_path = (SOURCE.parent / image_match.group(2)).resolve()
            next_index = index + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            caption = image_match.group(1)
            if next_index < len(lines) and lines[next_index].startswith("图 "):
                caption = lines[next_index].strip()
                next_index += 1
            compact_images = {"create_ctdp_task.png", "create_rsip_habit.png"}
            width = 3.7 if image_path.name in compact_images else 6.35
            add_figure(doc, image_path, caption, width)
            index = next_index
            continue

        match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if match:
            level = max(1, len(match.group(1)) - 1)
            p = doc.add_paragraph(match.group(2), style=f"Heading {level}")
            set_keep_with_next(p)
            index += 1
            continue

        if line.startswith("> "):
            add_note_box(doc, "重要提醒", line[2:].strip())
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|\s*[-:]+", lines[index + 1].strip()):
            headers = split_table_row(line)
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            col_count = len(headers)
            table = doc.add_table(rows=1, cols=col_count)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            widths = [TABLE_WIDTH_DXA // col_count] * col_count
            widths[-1] += TABLE_WIDTH_DXA - sum(widths)
            set_table_geometry(table, widths)
            for cell, text in zip(table.rows[0].cells, headers):
                set_cell_shading(cell, LIGHT_BLUE)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline_text(p, text, size=9.2)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            set_repeat_table_header(table.rows[0])
            for values in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, values):
                    p = cell.paragraphs[0]
                    add_inline_text(p, text, size=9.1)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline_text(p, numbered.group(2))
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_text(p, bullet.group(1))
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        p = doc.add_paragraph()
        add_inline_text(p, line)
        index += 1


def clear_doc_defaults(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def build_document() -> None:
    doc = Document()
    clear_doc_defaults(doc)
    configure_styles(doc)
    setup_page_furniture(doc)
    add_cover(doc)
    add_reading_map(doc)
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    if source_lines and source_lines[0].startswith("# "):
        source_lines = source_lines[1:]
    parse_markdown(doc, source_lines)

    properties = doc.core_properties
    properties.title = "个人工作台 Windows 版新手使用指导"
    properties.subject = "个人工作台 Windows 八栏功能逐项操作说明"
    properties.author = "个人工作台项目"
    properties.comments = "由项目现有功能说明和真实界面截图整理。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
